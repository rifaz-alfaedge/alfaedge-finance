import base64
import io
import json
from datetime import date
from unittest.mock import patch

import frappe
import openpyxl
from frappe.tests.utils import FrappeTestCase

from alfaedge_finance.alfaedge_finance.api.bank_statement_upload import upload_bank_statement
from alfaedge_finance.alfaedge_finance.bank_statement.matching import (
	BOOKED,
	CHECK,
	Candidate,
	LineInput,
	exact_sum_combinations,
	match_lines,
)
from alfaedge_finance.alfaedge_finance.bank_statement.narration import (
	counterparty_key,
	parse_narration,
	references_match,
)
from alfaedge_finance.alfaedge_finance.bank_statement.parsers import parse_statement_file
from alfaedge_finance.alfaedge_finance.bank_statement.suggestion import (
	SuggestionContext,
	find_invoice_combination,
	find_party_by_name,
	suggest,
)

COMPANY = "Code Dynamic Solutions Private Limited"
BANK_ACCOUNT = "Code Dynamic Solutions Private Limited - Axis Bank"
ACCOUNT_NO = "924020034470339"
EXPENSE_ACCOUNT = "Bank Charges - CDS"


def make_axis_statement(lines, opening=1000.0, account_no=ACCOUNT_NO, period=("01/09/2026", "03/09/2026")):
	"""Synthetic Axis "Statement Enquiry" workbook. lines: [(dd/mm/yyyy, particulars, amount, 'DR'|'CR')]."""
	workbook = openpyxl.Workbook()
	sheet = workbook.active
	sheet.append(["Name :- TEST COMPANY"])
	sheet.append(["Currency", None, "INR"])
	sheet.append(
		[f"Statement of Account No - {account_no} for the period (From : {period[0]} To : {period[1]} )"]
	)
	sheet.append([None])
	sheet.append(
		[
			"S.NO",
			"Transaction Date (dd/mm/yyyy)",
			"Value Date (dd/mm/yyyy)",
			"Particulars",
			"Amount(INR)",
			"Debit/Credit",
			"Balance(INR)",
			"Cheque Number",
			"Branch Name(SOL)",
		]
	)
	sheet.append(["1", "", "", "OPENING BALANCE", "", "", f"{opening:,.2f}", "", ""])
	balance = opening
	for index, (txn_date, particulars, amount, dr_cr) in enumerate(lines, start=2):
		balance = round(balance - amount if dr_cr == "DR" else balance + amount, 2)
		sheet.append(
			[str(index), txn_date, txn_date, particulars, f"{amount:,.2f}", dr_cr, f"{balance:,.2f}", "", ""]
		)
	sheet.append(["", "", "", "TRANSACTION TOTAL DR/CR", "", "", "", "", ""])
	sheet.append(["", "", "", "CLOSING BALANCE", "", "", f"{balance:,.2f}", "", ""])
	sheet.append(["Legend :"])
	buffer = io.BytesIO()
	workbook.save(buffer)
	return buffer.getvalue()


class TestAxisParser(FrappeTestCase):
	def test_parses_lines_period_and_balances(self):
		content = make_axis_statement(
			[
				("01/09/2026", "POS/FACEBK* ABC/+353/010926/10:00/611111111111", 1500.50, "DR"),
				("02/09/2026", "NEFT/HDFCH00000000001/ACME TRADERS/HDFC BANK/0001", 25000.00, "CR"),
			],
			opening=100000.0,
		)
		statement = parse_statement_file(content, "statement.xlsx")
		self.assertEqual(statement.account_number, ACCOUNT_NO)
		self.assertEqual(statement.from_date, date(2026, 9, 1))
		self.assertEqual(statement.to_date, date(2026, 9, 3))
		self.assertEqual(statement.currency, "INR")
		self.assertEqual(len(statement.lines), 2)
		self.assertEqual(statement.lines[0].withdrawal, 1500.50)
		self.assertEqual(statement.lines[1].deposit, 25000.00)
		self.assertEqual(statement.opening_balance, 100000.0)
		self.assertEqual(statement.closing_balance, 123499.50)

	def test_rejects_broken_running_balance(self):
		content = make_axis_statement([("01/09/2026", "POS/X/1/010926/10:00/611111111112", 100.0, "DR")])
		workbook = openpyxl.load_workbook(io.BytesIO(content))
		workbook.active.cell(row=7, column=7).value = "999.99"
		buffer = io.BytesIO()
		workbook.save(buffer)
		with self.assertRaises(frappe.ValidationError):
			parse_statement_file(buffer.getvalue(), "statement.xlsx")

	def test_rejects_unknown_layout(self):
		workbook = openpyxl.Workbook()
		workbook.active.append(["Date", "Narration", "Amount"])
		buffer = io.BytesIO()
		workbook.save(buffer)
		with self.assertRaises(frappe.ValidationError):
			parse_statement_file(buffer.getvalue(), "other.xlsx")


class TestNarration(FrappeTestCase):
	def test_channels(self):
		cases = {
			"POS/FACEBK* VRWBZF9G6/+35315530550/010426/23:00/609123083089": (
				"Card",
				"609123083089",
				"facebk",
			),
			"ECOM PUR/AMAZON PAY IN/1246624801/080426/05:50/609705127358": (
				"Card",
				"609705127358",
				"amazon pay in",
			),
			"NEFT/IN42609256645836/AL/TAZA SPARTA/ICICI BANK LIMITED/BULD69121796": (
				"NEFT",
				"IN42609256645836",
				"al taza sparta",
			),
			"NEFT/260001718903/14/AW0003809429/": ("Bulk Upload", "260001718903", "axis bulk upload"),
			"INB/NEFT/AXODH12041487058/C LOUNGE BUSINESS CEN/INDUSIND BANK//////": (
				"NEFT",
				"AXODH12041487058",
				"c lounge business cen",
			),
			"INB/950195056/GST TAX PAYMENT/": ("Net Banking", "950195056", "gst tax payment"),
			"IMPS/P2A/610610160765/ANSONCHI/FEDERALB/DYNAMICS/9199": ("IMPS", "610610160765", "ansonchi"),
			"BRN/REF NO.0081FIR2601568 USD 15160.80/RLZ": ("Branch", "0081FIR2601568", None),
			"VISA MERCH Refund /23/APR/26/GOOGLE CLOUD": ("Card Refund", None, "google cloud"),
			"Monthly Service Chrgs": ("Other", None, "monthly service chrgs"),
			"INB/IFT/Topcertifier/TPARTY TRANSFER": ("IFT", None, "topcertifier"),
		}
		for text, (channel, reference, key) in cases.items():
			narration = parse_narration(text)
			self.assertEqual(
				(narration.channel, narration.reference, counterparty_key(narration.counterparty)),
				(channel, reference, key),
				text,
			)

	def test_reference_match(self):
		self.assertTrue(references_match("IN42609256645836", "42609256645836"))
		self.assertTrue(references_match("0081FIR2601782", "0081FIR260178"))
		self.assertFalse(references_match("260001831033", "260001718903"))
		self.assertFalse(references_match("1234567", "12345678"))  # shorter side under 8 chars
		self.assertFalse(references_match(None, "12345678"))


class TestMatching(FrappeTestCase):
	def line(self, key, amount, reference=None, direction="Withdrawal", day=10):
		return LineInput(key, date(2026, 4, day), direction, amount, reference)

	def candidate(
		self, name, amount, reference=None, direction="Withdrawal", day=10, doctype="Payment Entry"
	):
		return Candidate(doctype, name, date(2026, 4, day), direction, amount, reference)

	def test_reference_then_amount_and_date(self):
		results = match_lines(
			[self.line("a", 761718.42, "0081FIR2601782", "Deposit"), self.line("b", 1910.0, None, "Deposit")],
			[
				self.candidate("PE-1", 761718.42, "0081FIR260178", "Deposit"),
				self.candidate("JE-1", 1910.0, "CONFIRMTKT RAIL", "Deposit", doctype="Journal Entry"),
			],
		)
		self.assertEqual((results["a"].status, results["a"].candidates[0].name), (BOOKED, "PE-1"))
		self.assertEqual(results["a"].basis, "Reference + amount")
		self.assertEqual((results["b"].status, results["b"].basis), (BOOKED, "Amount + date"))

	def test_conflicting_reference_is_check(self):
		results = match_lines(
			[self.line("a", 595500.0, "260001831033")],
			[self.candidate("JE-1", 595500.0, "260001718903", doctype="Journal Entry")],
		)
		self.assertEqual(results["a"].status, CHECK)

	def test_several_candidates_is_check(self):
		results = match_lines(
			[self.line("a", 500.0)],
			[self.candidate("PE-1", 500.0, day=9), self.candidate("PE-2", 500.0, day=12)],
		)
		self.assertEqual(results["a"].status, CHECK)
		self.assertEqual([[c.name for c in o] for o in results["a"].alternatives], [["PE-1"], ["PE-2"]])

	def test_voucher_never_claimed_twice_and_direction_amount_window_respected(self):
		results = match_lines(
			[self.line("a", 22140.0, "HDFCH1"), self.line("b", 22140.0, "HDFCH1")],
			[self.candidate("PE-1", 22140.0, "HDFCH1")],
		)
		self.assertEqual(len(results), 1)
		self.assertFalse(
			match_lines([self.line("a", 100.0)], [self.candidate("PE-1", 100.0, direction="Deposit")])
		)
		self.assertFalse(match_lines([self.line("a", 100.0)], [self.candidate("PE-1", 100.01)]))
		self.assertFalse(match_lines([self.line("a", 100.0, day=1)], [self.candidate("PE-1", 100.0, day=20)]))

	def test_several_vouchers_sharing_the_reference_book_the_line(self):
		results = match_lines(
			[self.line("a", 1000.0, "260001775107")],
			[
				self.candidate("PE-1", 600.0, "260001775107"),
				self.candidate("PE-2", 400.0, "260001775107"),
				self.candidate("PE-3", 400.0, "999"),
			],
		)
		self.assertEqual(results["a"].status, BOOKED)
		self.assertEqual(sorted(c.name for c in results["a"].candidates), ["PE-1", "PE-2"])

	def test_unique_exact_sum_is_check_and_ambiguous_sum_is_nothing(self):
		unique = match_lines(
			[self.line("a", 1000.0)],
			[self.candidate("PE-1", 700.0), self.candidate("PE-2", 300.0), self.candidate("PE-3", 450.0)],
		)
		self.assertEqual(unique["a"].status, CHECK)
		self.assertEqual(sorted(c.name for c in unique["a"].candidates), ["PE-1", "PE-2"])
		ambiguous = match_lines(
			[self.line("a", 1000.0)],
			[
				self.candidate("PE-1", 700.0),
				self.candidate("PE-2", 300.0),
				self.candidate("PE-3", 600.0),
				self.candidate("PE-4", 400.0),
			],
		)
		self.assertFalse(ambiguous)

	def test_exact_sum_combinations(self):
		self.assertEqual(exact_sum_combinations([1.1, 2.2, 3.3], 5.5, float, 3), [[2.2, 3.3]])
		self.assertEqual(exact_sum_combinations([1.0, 2.0], 5.0, float, 3), [])


class TestSuggestion(FrappeTestCase):
	def context(self, names=None):
		context = SuggestionContext(COMPANY, "924020034470339 - Axis Bank Thrippunithura - CDS", "INR")
		context._names = names or []
		context._history = {}
		return context

	def test_party_prefix_must_be_unique_and_long_enough(self):
		names = [
			("cloungebusinesscenterllp", "Supplier", "C Lounge Business Center LLP"),
			("acmeindustries", "Supplier", "Acme Industries"),
			("acmeindustriesexports", "Customer", "Acme Industries Exports"),
		]
		context = self.context(names)
		self.assertEqual(
			find_party_by_name(context, "C LOUNGE BUSINESS CEN", "Withdrawal"),
			("Supplier", "C Lounge Business Center LLP"),
		)
		self.assertIsNone(find_party_by_name(context, "C LOUNGE", "Withdrawal"))  # under 8 chars
		# Two parties share the prefix; the direction decides (money out -> supplier).
		self.assertEqual(
			find_party_by_name(context, "ACME INDUSTRIES", "Withdrawal"), ("Supplier", "Acme Industries")
		)

	def test_rule_wins_over_history_and_names(self):
		frappe.get_doc(
			{
				"doctype": "Bank Narration Rule",
				"direction": "Withdrawal",
				"counterparty_key": "Test Merchant ZZ",
				"suggested_doctype": "Journal Entry",
				"account": EXPENSE_ACCOUNT,
			}
		).insert()
		context = self.context()
		context._history = {("Withdrawal", "test merchant zz"): [("Journal Entry", None, None, "Other - X")]}
		result = suggest(context, "POS/TEST MERCHANT ZZ/KOCHI/010926/10:00/611111111113", "Withdrawal", 99.0)
		self.assertEqual(
			(result.suggested_doctype, result.account, result.confidence),
			("Journal Entry", EXPENSE_ACCOUNT, "High"),
		)

	def test_history_used_only_when_consistent(self):
		context = self.context()
		context._rules = {}
		context._history = {
			("Withdrawal", "same merchant"): [("Journal Entry", None, None, EXPENSE_ACCOUNT)] * 2,
			("Withdrawal", "mixed merchant"): [
				("Journal Entry", None, None, EXPENSE_ACCOUNT),
				("Journal Entry", None, None, "Other - X"),
			],
		}
		same = suggest(context, "POS/SAME MERCHANT/KOCHI/010926/10:00/611111111114", "Withdrawal", 50.0)
		mixed = suggest(context, "POS/MIXED MERCHANT/KOCHI/010926/10:00/611111111115", "Withdrawal", 50.0)
		self.assertEqual((same.account, same.confidence), (EXPENSE_ACCOUNT, "High"))
		self.assertEqual((mixed.account, mixed.confidence), (None, "Low"))

	def test_invoice_combination_split_plan(self):
		invoices = [
			frappe._dict(
				name="PI-A", supplier="Seller One", outstanding_amount=499.0, posting_date=date(2026, 9, 1)
			),
			frappe._dict(
				name="PI-B", supplier="Seller Two", outstanding_amount=1250.5, posting_date=date(2026, 9, 1)
			),
			frappe._dict(
				name="PI-C", supplier="Seller Two", outstanding_amount=300.0, posting_date=date(2026, 9, 2)
			),
		]
		with patch(
			"alfaedge_finance.alfaedge_finance.bank_statement.suggestion.frappe.get_all",
			return_value=invoices,
		):
			plan = find_invoice_combination(self.context(), "Withdrawal", 1749.5, date(2026, 9, 5))
		self.assertEqual(
			[(r["party"], r["against_name"], r["amount"]) for r in plan],
			[("Seller One", "PI-A", 499.0), ("Seller Two", "PI-B", 1250.5)],
		)

	def test_customer_tds_short_payment(self):
		from alfaedge_finance.alfaedge_finance.bank_statement.suggestion import find_invoice_with_tds

		invoices = [
			# 10% of net 29,000 = 2,900 withheld: 34,220 invoiced, 31,320 received.
			frappe._dict(
				name="SI-A",
				customer="Cust A",
				outstanding_amount=34220.0,
				grand_total=34220.0,
				net_total=29000.0,
			),
			# 10% of net 457,488 = 45,748.80, rounded to 45,749 by the customer.
			frappe._dict(
				name="SI-B",
				customer="Cust B",
				outstanding_amount=539835.84,
				grand_total=539835.84,
				net_total=457488.0,
			),
			# Part-paid: no TDS on the remainder.
			frappe._dict(
				name="SI-C",
				customer="Cust C",
				outstanding_amount=34220.0,
				grand_total=50000.0,
				net_total=29000.0,
			),
		]
		base = "alfaedge_finance.alfaedge_finance.bank_statement.suggestion"
		with (
			patch(f"{base}.frappe.get_all", return_value=invoices),
			patch(f"{base}.open_orders", return_value=[]),
		):
			self.assertEqual(
				find_invoice_with_tds(self.context(), 31320.0),
				[("Sales Invoice", "SI-A", "Cust A", 10, 2900.0)],
			)
			(hit,) = find_invoice_with_tds(self.context(), 494086.84)
			self.assertEqual((hit[1], hit[3], hit[4]), ("SI-B", 10, 45749.0))
			self.assertEqual(find_invoice_with_tds(self.context(), 30000.0), [])

		# A Sales Order advance: 2% of net 3,100 = 62 withheld from 3,658.
		order = {"doctype": "Sales Order", "name": "SO-1", "party": "Cust D", "amount": 3658.0}
		with (
			patch(f"{base}.frappe.get_all", return_value=[]),
			patch(f"{base}.open_orders", return_value=[order]),
			patch(f"{base}.frappe.db.get_value", return_value=frappe._dict(net_total=3100.0, advance_paid=0)),
		):
			self.assertEqual(
				find_invoice_with_tds(self.context(), 3596.0), [("Sales Order", "SO-1", "Cust D", 2, 62.0)]
			)

	def test_open_orders_are_settleable(self):
		from alfaedge_finance.alfaedge_finance.bank_statement.suggestion import find_open_documents

		order = {
			"doctype": "Purchase Order",
			"name": "PO-1",
			"party": "Supp X",
			"amount": 7399.0,
			"date": None,
		}
		base = "alfaedge_finance.alfaedge_finance.bank_statement.suggestion"
		with (
			patch(f"{base}.frappe.get_all", return_value=[]),
			patch(f"{base}.frappe.db.sql", return_value=[]),
			patch(
				f"{base}.open_orders",
				side_effect=lambda doctype, *a, **k: [order] if doctype == "Purchase Order" else [],
			),
		):
			self.assertEqual(
				find_open_documents(self.context(), "Withdrawal", 7399.0),
				[("Purchase Order", "PO-1", "Supplier", "Supp X")],
			)
			self.assertEqual(find_open_documents(self.context(), "Deposit", 7399.0), [])

	def test_party_named_in_narration(self):
		from alfaedge_finance.alfaedge_finance.bank_statement.narration import parse_narration
		from alfaedge_finance.alfaedge_finance.bank_statement.suggestion import party_named_in_narration

		def named(party, text):
			return party_named_in_narration(party, text, parse_narration(text).counterparty)

		self.assertTrue(named("OpenAI OpCo, LLC", "ECOM PUR/OPENAI/+14158799686/210926/05:23/626405671138"))
		self.assertTrue(
			named("Anthropic, PBC", "ECOM PUR/ANTHROPIC* CL/+14152360599/090926/09:49/625209503629")
		)
		self.assertTrue(
			named(
				"ZenLeads Inc. (dba Apollo.io)", "ECOM PUR/APOLLO.IO/+14156912009/170926/04:55/626004039089"
			)
		)
		self.assertTrue(
			named(
				"Akamai Technologies International AG", "ECOM PUR/LINODE . AKAM/CAMBRIDGE/080926/03:37/6250"
			)
		)
		self.assertTrue(named("Exa Labs Inc.", "ECOM PUR/EXA.AI/+18316182513/170926/05:01/626005110363"))
		self.assertFalse(
			named("OpenAI OpCo, LLC", "ECOM PUR/ANTHROPIC* CL/+14152360599/090926/09:49/625209503629")
		)
		self.assertFalse(
			named("Google Cloud India Private Limited", "ECOM PUR/WWW AMAZON IN/1243054000/010926/13:42/6244")
		)

	def test_foreign_invoice_by_name_and_rate(self):
		from alfaedge_finance.alfaedge_finance.bank_statement.suggestion import find_foreign_documents

		invoices = [
			frappe._dict(
				name="PI-USD",
				supplier="OpenAI OpCo, LLC",
				supplier_name="OpenAI OpCo, LLC",
				currency="USD",
				outstanding_amount=50.0,
				conversion_rate=95.0,
			)
		]
		base = "alfaedge_finance.alfaedge_finance.bank_statement.suggestion"
		narration = "ECOM PUR/OPENAI/+14158799686/210926/05:23/626405671138"
		with (
			patch(f"{base}.frappe.get_all", return_value=invoices),
			patch(f"{base}.get_exchange_rate", return_value=95.82),
		):
			(hit,) = find_foreign_documents(self.context(), "Withdrawal", 4995.11, "2026-09-21", narration)
			self.assertEqual(hit[:4], ("Purchase Invoice", "PI-USD", "Supplier", "OpenAI OpCo, LLC"))
			# Way off the market rate - not this invoice.
			self.assertEqual(
				find_foreign_documents(self.context(), "Withdrawal", 9000.0, "2026-09-21", narration), []
			)
			# Name not in the narration and no known party - no guess.
			self.assertEqual(
				find_foreign_documents(
					self.context(), "Withdrawal", 4995.11, "2026-09-21", "ECOM PUR/SOMETHING/1/2"
				),
				[],
			)
			# A narration quoting the foreign amount matches exactly, name or not.
			(quoted,) = find_foreign_documents(
				self.context(), "Withdrawal", 5500.0, "2026-09-21", "BRN/REF NO.0081FIR26 USD 50.00/RLZ"
			)
			self.assertEqual(quoted[1], "PI-USD")


class TestBankStatementReview(FrappeTestCase):
	"""End to end against the real company bank account, with dates/amounts chosen not to
	collide with real data. Everything is rolled back by FrappeTestCase."""

	REF = "611199990001"

	def setUp(self):
		frappe.set_user("Administrator")
		# Uploads auto-process by default; these tests drive each step themselves.
		self.auto = patch(
			"alfaedge_finance.alfaedge_finance.doctype.bank_statement_review.bank_statement_review.BankStatementReview.auto_process",
			return_value={"drafted": 0, "queued": 0, "failed": []},
		)
		self.auto.start()
		self.addCleanup(self.auto.stop)
		self.gl_account = frappe.db.get_value("Bank Account", BANK_ACCOUNT, "account")
		je = frappe.get_doc(
			{
				"doctype": "Journal Entry",
				"voucher_type": "Bank Entry",
				"company": COMPANY,
				"posting_date": "2026-09-01",
				"cheque_no": self.REF,
				"cheque_date": "2026-09-01",
				"accounts": [
					{"account": self.gl_account, "credit_in_account_currency": 4321.87},
					{"account": EXPENSE_ACCOUNT, "debit_in_account_currency": 4321.87},
				],
			}
		).insert()
		je.submit()
		self.booked_je = je.name

	def tearDown(self):
		# FrappeTestCase only rolls back per class; each test here uploads the same statement.
		frappe.db.rollback()

	def upload(self):
		content = make_axis_statement(
			[
				("01/09/2026", f"POS/TEST SHOP ONE/KOCHI/010926/10:00/{self.REF}", 4321.87, "DR"),
				("02/09/2026", "POS/UNKNOWN MERCHANT QQ/KOCHI/020926/11:00/611199990002", 1234.56, "DR"),
			]
		)
		name = upload_bank_statement("test.xlsx", base64.b64encode(content).decode())
		return frappe.get_doc("Bank Statement Review", name)

	def test_full_flow(self):
		review = self.upload()
		self.assertEqual((str(review.from_date), str(review.to_date)), ("2026-09-01", "2026-09-03"))
		booked, missing = review.lines
		self.assertEqual(booked.line_status, "Already Booked")
		self.assertEqual([(v.line_no, v.voucher_no) for v in review.vouchers], [(1, self.booked_je)])
		self.assertEqual(missing.line_status, "Suggested")

		# Create the missing record as a draft - never submitted by the tool.
		(voucher,) = review.create_draft(
			missing.name, {"suggested_doctype": "Journal Entry", "account": EXPENSE_ACCOUNT}
		)
		self.assertEqual(frappe.db.get_value("Journal Entry", voucher["name"], "docstatus"), 0)
		self.assertTrue(
			frappe.db.exists("Bank Narration Rule", {"rule_key": "Withdrawal|unknown merchant qq"})
		)

		# Submitting the draft books the line straight away (voucher_sync), no re-run needed.
		frappe.get_doc("Journal Entry", voucher["name"]).submit()
		review = frappe.get_doc("Bank Statement Review", review.name)
		self.assertEqual(review.lines[1].line_status, "Already Booked")
		self.assertEqual(review.rerun()["counts"], {"Already Booked": 2})
		self.assertEqual(review.status, "Reviewed")

		sent = review.send_to_bank_reconciliation()
		self.assertEqual(sent["created"], 2)
		self.assertEqual(review.status, "Sent")
		transaction = frappe.get_doc("Bank Transaction", review.lines[0].bank_transaction)
		self.assertEqual((transaction.docstatus, transaction.withdrawal), (1, 4321.87))

		# The same statement again: everything is already imported and nothing is re-sent.
		again = self.upload()
		self.assertTrue(all(line.line_status == "Already Imported" for line in again.lines))
		self.assertEqual(again.send_to_bank_reconciliation(), {"created": 0, "linked": 2})

	def test_create_high_confidence_drafts_in_bulk(self):
		frappe.get_doc(
			{
				"doctype": "Bank Narration Rule",
				"direction": "Withdrawal",
				"counterparty_key": "unknown merchant qq",
				"suggested_doctype": "Journal Entry",
				"account": EXPENSE_ACCOUNT,
			}
		).insert()
		review = self.upload()
		self.assertEqual(review.lines[1].confidence, "High")
		result = review.create_high_confidence_drafts()
		self.assertEqual((len(result["created"]), result["failed"]), (1, []))
		self.assertEqual(review.lines[1].line_status, "Draft Created")
		self.assertEqual(frappe.db.get_value("Journal Entry", result["created"][0], "docstatus"), 0)

	def test_bulk_drafts_skip_a_failing_line_and_keep_going(self):
		frappe.get_doc(
			{
				"doctype": "Bank Narration Rule",
				"direction": "Withdrawal",
				"counterparty_key": "unknown merchant qq",
				"suggested_doctype": "Journal Entry",
				"account": EXPENSE_ACCOUNT,
			}
		).insert()
		review = self.upload()
		line = review.lines[1]
		with patch(
			"alfaedge_finance.alfaedge_finance.doctype.bank_statement_review.bank_statement_review.build_voucher",
			side_effect=frappe.ValidationError("cannot build this one"),
		):
			result = review.create_high_confidence_drafts()
		self.assertEqual(len(result["created"]), 0)
		self.assertEqual(len(result["failed"]), 1)
		self.assertIn("cannot build this one", result["failed"][0])
		self.assertEqual((line.line_status, len(review.vouchers)), ("Suggested", 1))

	def test_settleable_search_shows_amounts(self):
		from alfaedge_finance.alfaedge_finance.bank_statement.link_queries import search_settleable

		for doctype in (
			"Sales Invoice",
			"Purchase Invoice",
			"Sales Order",
			"Purchase Order",
			"Expense Claim",
		):
			for name, _party, _date, amount in search_settleable(
				doctype, "", "name", 0, 5, {"company": COMPANY}
			):
				self.assertTrue(amount.startswith("Outstanding "), (doctype, name, amount))
		self.assertEqual(search_settleable("Journal Entry", "", "name", 0, 5, {}), [])

	def test_auto_process_drafts_medium_and_submits_high(self):
		from alfaedge_finance.alfaedge_finance.doctype.bank_statement_review import bank_statement_review

		self.auto.stop()
		frappe.get_doc(
			{
				"doctype": "Bank Narration Rule",
				"direction": "Withdrawal",
				"counterparty_key": "unknown merchant qq",
				"suggested_doctype": "Journal Entry",
				"account": EXPENSE_ACCOUNT,
			}
		).insert()
		with (
			patch.object(bank_statement_review.frappe, "enqueue") as enqueue,
			patch.object(bank_statement_review.BankStatementReview, "_remember_rule") as remember,
		):
			review = self.upload()
		remember.assert_not_called()  # automatic drafts never teach rules
		line = review.lines[1]
		self.assertEqual((line.confidence, line.line_status), ("High", "Draft Created"))
		(queued,) = enqueue.call_args.kwargs["vouchers"]
		self.assertEqual(queued, ("Journal Entry", review.vouchers[-1].voucher_no))
		self.assertIn("1 queued for submission", review.last_run_summary)

		# Deleting what was auto-created: the suggestion stays, but Re-run doesn't redo it.
		review.delete_draft("Journal Entry", review.vouchers[-1].voucher_no)
		self.assertEqual((line.line_status, line.confidence), ("Suggested", None))
		with patch.object(bank_statement_review.frappe, "enqueue") as enqueue:
			review.rerun()
		enqueue.assert_not_called()
		self.assertEqual((line.line_status, line.account), ("Suggested", EXPENSE_ACCOUNT))

		# Switched off: nothing is created.
		review.auto_submit_high = review.auto_draft_medium = 0
		self.assertEqual(review.auto_process(), {"drafted": 0, "queued": 0, "failed": []})

	def test_older_review_keeps_its_lines_when_a_newer_one_overlaps(self):
		older = self.upload()
		older.create_draft(
			older.lines[1].name, {"suggested_doctype": "Journal Entry", "account": EXPENSE_ACCOUNT}
		)
		newer = self.upload()
		self.assertTrue(all(line.line_status == "Already Imported" for line in newer.lines))
		older = frappe.get_doc("Bank Statement Review", older.name)
		older.rerun()
		self.assertEqual([line.line_status for line in older.lines], ["Already Booked", "Draft Created"])

	def test_same_amount_documents_closest_date_wins(self):
		from alfaedge_finance.alfaedge_finance.bank_statement.suggestion import (
			SuggestionContext,
			find_open_documents,
		)

		context = SuggestionContext(COMPANY, "x", "INR")
		rows = [
			frappe._dict(name="PI-JUNE", supplier="S", posting_date=date(2026, 6, 2)),
			frappe._dict(name="PI-SEPT", supplier="S", posting_date=date(2026, 9, 1)),
			frappe._dict(name="PI-LATER", supplier="S", posting_date=date(2026, 9, 5)),
		]
		base = "alfaedge_finance.alfaedge_finance.bank_statement.suggestion"
		with (
			patch(f"{base}.frappe.get_all", return_value=rows),
			patch(f"{base}.frappe.db.sql", return_value=[]),
			patch(f"{base}.open_orders", return_value=[]),
		):
			found = find_open_documents(context, "Withdrawal", 19969.2, "Supplier", "S", date(2026, 9, 1))
		self.assertEqual([f[1] for f in found], ["PI-SEPT", "PI-JUNE", "PI-LATER"])

	def _two_unmatched_lines(self):
		content = make_axis_statement(
			[
				("01/09/2026", "POS/PREPAID THING ZZ/KOCHI/010926/10:00/611199990011", 5000.0, "DR"),
				("02/09/2026", "POS/ANOTHER THING ZZ/KOCHI/020926/11:00/611199990012", 700.0, "DR"),
			]
		)
		name = upload_bank_statement("adv.xlsx", base64.b64encode(content).decode())
		return frappe.get_doc("Bank Statement Review", name)

	def test_mark_single_line_as_advance(self):
		supplier = frappe.get_all("Supplier", filters={"disabled": 0}, pluck="name", limit=1)[0]
		review = self._two_unmatched_lines()
		line = review.lines[0]
		result = review.mark_as_advance(
			[line.name], {"party_type": "Supplier", "party": supplier, "reference": "PRO-123", "remember": 1}
		)
		self.assertEqual((len(result["created"]), result["failed"]), (1, []))
		pe = frappe.get_doc("Payment Entry", result["created"][0])
		self.assertEqual(
			(pe.docstatus, pe.party, len(pe.references), pe.unallocated_amount), (0, supplier, 0, 5000.0)
		)
		self.assertIn("PRO-123", pe.remarks)
		self.assertTrue(line.is_advance)
		self.assertEqual(review.vouchers[-1].is_advance, 1)
		rule = frappe.get_doc("Bank Narration Rule", {"rule_key": "Withdrawal|prepaid thing zz"})
		self.assertEqual((rule.party, rule.treat_as_advance), (supplier, 1))

		# The remembered payee is suggested as an advance next time.
		from alfaedge_finance.alfaedge_finance.bank_statement.suggestion import SuggestionContext, suggest

		again = suggest(
			SuggestionContext(COMPANY, "x", "INR"),
			"POS/PREPAID THING ZZ/KOCHI/150926/10:00/611199990019",
			"Withdrawal",
			5000.0,
			date(2026, 9, 15),
		)
		self.assertEqual(
			(again.is_advance, again.party, again.against_name, again.confidence),
			(True, supplier, None, "High"),
		)

	def test_mark_lines_as_advance_in_bulk_and_employee_account(self):
		review = self._two_unmatched_lines()
		# No party for either line and none given: both fail, nothing is created.
		result = review.mark_as_advance([line.name for line in review.lines], {"party_type": "Supplier"})
		self.assertEqual((result["created"], len(result["failed"])), ([], 2))

		review = frappe.get_doc("Bank Statement Review", review.name)
		employee = frappe.get_all("Employee", filters={"status": "Active"}, pluck="name", limit=1)[0]
		result = review.mark_as_advance(
			[line.name for line in review.lines], {"party_type": "Employee", "party": employee}
		)
		self.assertEqual(len(result["created"]), 2)
		advance_account = frappe.get_cached_value("Company", COMPANY, "default_employee_advance_account")
		for name in result["created"]:
			self.assertEqual(frappe.db.get_value("Payment Entry", name, "paid_to"), advance_account)

	def test_rejects_statement_for_another_account(self):
		content = make_axis_statement(
			[("01/09/2026", "POS/X/1/010926/10:00/611199990003", 10.0, "DR")], account_no="111122223333"
		)
		with self.assertRaises(frappe.ValidationError):
			upload_bank_statement("test.xlsx", base64.b64encode(content).decode(), bank_account=BANK_ACCOUNT)

	def test_delete_drafts_from_review_and_from_their_own_form(self):
		review = self.upload()
		missing = review.lines[1]
		values = {"suggested_doctype": "Journal Entry", "account": EXPENSE_ACCOUNT}

		(voucher,) = review.create_draft(missing.name, values)
		self.assertEqual(missing.line_status, "Draft Created")
		review.delete_draft("Journal Entry", voucher["name"])
		self.assertFalse(frappe.db.exists("Journal Entry", voucher["name"]))
		self.assertEqual((missing.line_status, len(review.vouchers)), ("Suggested", 1))

		(voucher,) = review.create_draft(missing.name, values)
		frappe.delete_doc("Journal Entry", voucher["name"])  # link from the review doesn't block it
		review = frappe.get_doc("Bank Statement Review", review.name)
		self.assertEqual((review.lines[1].line_status, len(review.vouchers)), ("Suggested", 1))

	def test_submit_draft_from_the_review(self):
		review = self.upload()
		(voucher,) = review.create_draft(
			review.lines[1].name, {"suggested_doctype": "Journal Entry", "account": EXPENSE_ACCOUNT}
		)
		review.submit_draft("Journal Entry", voucher["name"])
		self.assertEqual(frappe.db.get_value("Journal Entry", voucher["name"], "docstatus"), 1)
		self.assertEqual(review.lines[1].line_status, "Already Booked")
		self.assertEqual(review.vouchers[-1].voucher_status, "Submitted")
		with self.assertRaises(frappe.ValidationError):  # matched, not created here
			review.submit_draft("Journal Entry", self.booked_je)

	def test_bulk_submit_is_queued_and_the_job_submits_each_draft(self):
		from alfaedge_finance.alfaedge_finance.doctype.bank_statement_review import bank_statement_review

		review = self.upload()
		(voucher,) = review.create_draft(
			review.lines[1].name, {"suggested_doctype": "Journal Entry", "account": EXPENSE_ACCOUNT}
		)
		selected = [{"voucher_type": "Journal Entry", "voucher_no": voucher["name"]}]
		with patch.object(bank_statement_review.frappe, "enqueue") as enqueue:
			self.assertEqual(review.submit_drafts_in_background(selected), {"queued": 1})
		self.assertEqual(enqueue.call_args.kwargs["vouchers"], [("Journal Entry", voucher["name"])])
		self.assertEqual(frappe.db.get_value("Journal Entry", voucher["name"], "docstatus"), 0)
		with self.assertRaises(frappe.ValidationError):  # only drafts created from this review
			review.submit_drafts_in_background(
				[{"voucher_type": "Journal Entry", "voucher_no": self.booked_je}]
			)

		# The job itself (run inline; its per-document commits are stubbed out here).
		with patch.object(frappe.db, "commit"), patch.object(frappe, "publish_realtime") as realtime:
			bank_statement_review.submit_drafts_job(
				review.name, [("Journal Entry", voucher["name"]), ("Journal Entry", "NO-SUCH-JE")]
			)
		self.assertEqual(frappe.db.get_value("Journal Entry", voucher["name"], "docstatus"), 1)
		done = realtime.call_args.args[1]
		self.assertEqual((done["submitted"], len(done["failed"])), ([voucher["name"]], 1))
		review = frappe.get_doc("Bank Statement Review", review.name)
		self.assertEqual(review.lines[1].line_status, "Already Booked")

	def test_edit_reference_of_a_draft_only(self):
		review = self.upload()
		self.assertEqual(review.vouchers[0].reference_no, self.REF)  # the matched JE's cheque no
		(voucher,) = review.create_draft(
			review.lines[1].name, {"suggested_doctype": "Journal Entry", "account": EXPENSE_ACCOUNT}
		)
		review.update_voucher_reference("Journal Entry", voucher["name"], " UTR-CHANGED-1 ")
		self.assertEqual(frappe.db.get_value("Journal Entry", voucher["name"], "cheque_no"), "UTR-CHANGED-1")
		self.assertEqual(review.vouchers[-1].reference_no, "UTR-CHANGED-1")
		with self.assertRaises(frappe.ValidationError):  # submitted: ERPNext doesn't allow it
			review.update_voucher_reference("Journal Entry", self.booked_je, "X")

	def test_cancelling_a_booked_voucher_unbooks_the_line(self):
		review = self.upload()
		frappe.get_doc("Journal Entry", self.booked_je).cancel()
		review = frappe.get_doc("Bank Statement Review", review.name)
		self.assertEqual((review.lines[0].line_status, len(review.vouchers)), ("Suggested", 0))

	def test_review_delete_takes_its_drafts_until_sent(self):
		review = self.upload()
		(draft,) = review.create_draft(
			review.lines[1].name, {"suggested_doctype": "Journal Entry", "account": EXPENSE_ACCOUNT}
		)
		frappe.delete_doc("Bank Statement Review", review.name)
		self.assertFalse(frappe.db.exists("Journal Entry", draft["name"]))
		self.assertEqual(frappe.db.get_value("Journal Entry", self.booked_je, "docstatus"), 1)

		review = self.upload()
		review.send_to_bank_reconciliation()
		with self.assertRaises(frappe.ValidationError):
			frappe.delete_doc("Bank Statement Review", review.name)

	def test_split_payment_creates_one_draft_per_party(self):
		review = self.upload()
		line = review.lines[1]
		suppliers = frappe.get_all("Supplier", filters={"disabled": 0}, pluck="name", limit=2)
		plan = [
			{"party_type": "Supplier", "party": suppliers[0], "amount": 1000.0},
			{"party_type": "Supplier", "party": suppliers[1], "amount": 200.0},
			{"party_type": "Supplier", "party": suppliers[1], "amount": 34.56},
		]
		with self.assertRaises(frappe.ValidationError):
			review.create_draft(line.name, {"split_plan": plan[:2], "use_split": True})  # doesn't add up
		review = frappe.get_doc("Bank Statement Review", review.name)
		vouchers = review.create_draft(review.lines[1].name, {"split_plan": plan, "use_split": True})
		self.assertEqual(len(vouchers), 2)
		amounts = sorted(frappe.db.get_value("Payment Entry", v["name"], "paid_amount") for v in vouchers)
		self.assertEqual(amounts, [234.56, 1000.0])
		self.assertTrue(
			all(frappe.db.get_value("Payment Entry", v["name"], "docstatus") == 0 for v in vouchers)
		)
		self.assertEqual(review.lines[1].line_status, "Draft Created")


def make_bulk_file(rows, debit_account=ACCOUNT_NO):
	"""Axis bulk-upload file as the Bulk Payment CSV page writes it: rows of
	(beneficiary name, account no, amount)."""
	workbook = openpyxl.Workbook()
	sheet = workbook.active
	sheet.append(
		[
			"Debit Account Number (Mandatory)",
			"Transaction Amount (Mandatory)",
			"Transaction Currency (Mandatory)",
			"Beneficiary Name (Mandatory)",
			"Beneficiary Account Number (Mandatory)",
			"Beneficiary IFSC Code (Mandatory)",
			"Transaction Date (Mandatory)",
		]
	)
	for name, account_no, amount in rows:
		sheet.append([debit_account, amount, "INR", name, account_no, "UTIB0000001", "02/09/2026"])
	buffer = io.BytesIO()
	workbook.save(buffer)
	return buffer.getvalue()


class TestBulkPayments(FrappeTestCase):
	EMPLOYEE_ACCOUNT = "999900001111"
	SUPPLIER_ACCOUNT = "999900002222"

	def setUp(self):
		frappe.set_user("Administrator")
		auto = patch(
			"alfaedge_finance.alfaedge_finance.doctype.bank_statement_review.bank_statement_review.BankStatementReview.auto_process",
			return_value={"drafted": 0, "queued": 0, "failed": []},
		)
		auto.start()
		self.addCleanup(auto.stop)
		self.employee = frappe.get_all("Employee", filters={"status": "Active"}, pluck="name", limit=1)[0]
		self.supplier = frappe.get_all("Supplier", filters={"disabled": 0}, pluck="name", limit=1)[0]
		bank = frappe.get_all("Bank", pluck="name", limit=1)[0]
		for party_type, party, number in (
			("Employee", self.employee, self.EMPLOYEE_ACCOUNT),
			("Supplier", self.supplier, self.SUPPLIER_ACCOUNT),
		):
			frappe.get_doc(
				{
					"doctype": "Bank Account",
					"account_name": f"BSR Test {number}",
					"bank": bank,
					"party_type": party_type,
					"party": party,
					"bank_account_no": f" {number} ",  # stored numbers can carry stray spaces
				}
			).insert()

	def tearDown(self):
		frappe.db.rollback()

	def test_parse_bulk_file(self):
		from alfaedge_finance.alfaedge_finance.bank_statement.bulk_upload import parse_bulk_file

		rows = parse_bulk_file(
			make_bulk_file([("A", 999900001111.0, 100.5), ("B", "9999 0000 2222", "1,000")]), "bulk.xlsx"
		)
		self.assertEqual(
			[(r.account_no, r.amount) for r in rows], [("999900001111", 100.5), ("999900002222", 1000.0)]
		)
		self.assertEqual(rows[0].debit_account_no, ACCOUNT_NO)

	def test_resolve_rows(self):
		from alfaedge_finance.alfaedge_finance.bank_statement import bulk_matching
		from alfaedge_finance.alfaedge_finance.bank_statement.bulk_upload import BulkRow

		def row(no, account, amount):
			return BulkRow(no, f"Row {no}", account, "", amount, ACCOUNT_NO, "", "")

		claims = [
			{"doctype": "Expense Claim", "name": "EC-1", "amount": 300.0},
			{"doctype": "Expense Claim", "name": "EC-2", "amount": 200.0},
			{"doctype": "Expense Claim", "name": "EC-3", "amount": 999.0},
		]
		base = "alfaedge_finance.alfaedge_finance.bank_statement.bulk_matching"
		with (
			patch(f"{base}.open_expense_claims", return_value=claims),
			patch(f"{base}.unpaid_salary_slip", return_value=None),
			patch(f"{base}.open_purchase_invoices", return_value=[]),
		):
			plan = bulk_matching.resolve_rows(
				COMPANY,
				"INR",
				"2026-09-02",
				[
					row(1, self.EMPLOYEE_ACCOUNT, 500.0),
					row(2, self.SUPPLIER_ACCOUNT, 75.0),
					row(3, "123123123", 10.0),
				],
			)
		self.assertEqual(
			[(p["status"], p["party"], [d["name"] for d in p["documents"]]) for p in plan],
			[
				("Matched", self.employee, ["EC-1", "EC-2"]),
				("On-account", self.supplier, []),
				("Unresolved", None, []),
			],
		)
		self.assertEqual(bulk_matching.summarise(plan)[1], "Low")

		order = {"doctype": "Purchase Order", "name": "PO-9", "party": self.supplier, "amount": 75.0}
		with (
			patch(f"{base}.open_purchase_invoices", return_value=[]),
			patch(f"{base}.open_orders", return_value=[order]),
		):
			(advance,) = bulk_matching.resolve_rows(
				COMPANY, "INR", "2026-09-02", [row(1, self.SUPPLIER_ACCOUNT, 75.0)]
			)
		self.assertEqual((advance["status"], advance["kind"]), ("Matched", "purchase_order"))

	def test_oldest_first_allocation_up_to_file_date(self):
		from alfaedge_finance.alfaedge_finance.bank_statement import bulk_matching
		from alfaedge_finance.alfaedge_finance.bank_statement.bulk_upload import BulkRow

		claims = [
			{"doctype": "Expense Claim", "name": "EC-OLD", "amount": 1000.0, "date": date(2026, 8, 1)},
			{"doctype": "Expense Claim", "name": "EC-MID", "amount": 700.0, "date": date(2026, 9, 1)},
			{"doctype": "Expense Claim", "name": "EC-LATE", "amount": 628.16, "date": date(2026, 9, 23)},
		]
		base = "alfaedge_finance.alfaedge_finance.bank_statement.bulk_matching"

		def resolve(amount):
			row = BulkRow(1, "A", self.EMPLOYEE_ACCOUNT, "", amount, ACCOUNT_NO, "", "12/09/2026")
			with (
				patch(f"{base}.open_expense_claims", return_value=claims),
				patch(f"{base}.unpaid_salary_slip", return_value=None),
			):
				return bulk_matching.resolve_rows(COMPANY, "INR", "2026-09-14", [row])[0]

		# More than the claims up to 12-09 (the late one is excluded): all of them, rest on account.
		over = resolve(1975.23)
		self.assertEqual(
			[(d["name"], d["amount"]) for d in over["documents"]], [("EC-OLD", 1000.0), ("EC-MID", 700.0)]
		)
		self.assertEqual(over["status"], "On-account")
		self.assertIn("275.23", over["note"])
		# Less than the claims: oldest first, the last one partly.
		under = resolve(1200.0)
		self.assertEqual(
			[(d["name"], d["amount"]) for d in under["documents"]], [("EC-OLD", 1000.0), ("EC-MID", 200.0)]
		)
		self.assertEqual(under["status"], "Matched")

	def _review_with_bulk_line(self, amount):
		content = make_axis_statement(
			[("02/09/2026", "NEFT/611199990009/2/AW0009999999////", amount, "DR")],
			period=("02/09/2026", "02/09/2026"),
		)
		name = upload_bank_statement("test.xlsx", base64.b64encode(content).decode())
		review = frappe.get_doc("Bank Statement Review", name)
		self.assertEqual((review.lines[0].channel, review.lines[0].bulk_count), ("Bulk Upload", 2))
		return review

	def test_bulk_line_is_never_guessed(self):
		# A booked voucher with the same amount and date but no batch reference must not match.
		gl_account = frappe.db.get_value("Bank Account", BANK_ACCOUNT, "account")

		def journal_entry(cheque_no):
			je = frappe.get_doc(
				{
					"doctype": "Journal Entry",
					"voucher_type": "Bank Entry",
					"company": COMPANY,
					"posting_date": "2026-09-02",
					"cheque_no": cheque_no,
					"cheque_date": "2026-09-02",
					"accounts": [
						{"account": gl_account, "credit_in_account_currency": 777.77},
						{"account": EXPENSE_ACCOUNT, "debit_in_account_currency": 777.77},
					],
				}
			).insert()
			je.submit()
			return je.name

		journal_entry("SOMETHING-ELSE-1")
		review = self._review_with_bulk_line(777.77)
		line = review.lines[0]
		self.assertEqual((line.line_status, line.suggested_doctype, line.party), ("Suggested", None, None))
		self.assertIn("upload the bank's bulk payment file", line.suggestion_source)
		# Not enforced: the reviewer can still book it by hand (e.g. a single payment sent as
		# a bulk upload), the choice survives a re-run, and no rule is learned from it.
		(voucher,) = review.create_draft(
			line.name, {"suggested_doctype": "Journal Entry", "account": EXPENSE_ACCOUNT}
		)
		review.delete_draft("Journal Entry", voucher["name"])
		review.rerun()
		self.assertEqual((line.suggested_doctype, line.account), ("Journal Entry", EXPENSE_ACCOUNT))
		self.assertFalse(frappe.db.exists("Bank Narration Rule", {"counterparty_key": "axis bulk upload"}))
		with self.assertRaises(frappe.ValidationError):
			frappe.get_doc(
				{
					"doctype": "Bank Narration Rule",
					"direction": "Withdrawal",
					"counterparty_key": "AXIS BULK UPLOAD",
					"suggested_doctype": "Journal Entry",
					"account": EXPENSE_ACCOUNT,
				}
			).insert()

		# ...but a voucher carrying the bank's batch reference is the booking.
		frappe.db.rollback()
		booked = journal_entry("611199990009")
		review = self._review_with_bulk_line(777.77)
		self.assertEqual(review.lines[0].line_status, "Already Booked")
		self.assertEqual(review.vouchers[0].voucher_no, booked)

	def test_upload_checks_edit_and_create_journal_entry(self):
		review = self._review_with_bulk_line(1234.5)
		line = review.lines[0]

		def upload(rows, **kwargs):
			return review.upload_bulk_file(
				line.name, "bulk.xlsx", base64.b64encode(make_bulk_file(rows, **kwargs)).decode()
			)

		with self.assertRaises(frappe.ValidationError):  # doesn't add up to the bank line
			upload([("A", self.SUPPLIER_ACCOUNT, 1000.0)])
		with self.assertRaises(frappe.ValidationError):  # debits a different account
			upload([("A", self.SUPPLIER_ACCOUNT, 1234.5)], debit_account="111122223333")

		review = frappe.get_doc("Bank Statement Review", review.name)
		line = review.lines[0]
		with patch(
			"alfaedge_finance.alfaedge_finance.bank_statement.bulk_matching.open_purchase_invoices",
			return_value=[],
		):
			result = upload([("A", self.SUPPLIER_ACCOUNT, 1000.0), ("B", "555", 234.5)])
		self.assertEqual([r["status"] for r in result["plan"]], ["On-account", "Unresolved"])
		with self.assertRaises(frappe.ValidationError):  # an Unresolved row blocks creation
			review.create_draft(line.name)

		review = frappe.get_doc("Bank Statement Review", review.name)
		line = review.lines[0]
		plan = json.loads(line.bulk_plan)
		plan[1].update(party_type="Employee", party=self.employee)
		review.save_bulk_plan(line.name, plan)
		self.assertEqual(
			[r["status"] for r in json.loads(review.lines[0].bulk_plan)], ["On-account", "On-account"]
		)

		(voucher,) = review.create_draft(review.lines[0].name)
		je = frappe.get_doc("Journal Entry", voucher["name"])
		self.assertEqual((je.docstatus, je.cheque_no, je.total_credit), (0, "611199990009", 1234.5))
		self.assertEqual(
			sorted((a.party, a.debit) for a in je.accounts if a.debit),
			sorted([(self.supplier, 1000.0), (self.employee, 234.5)]),
		)
		self.assertEqual(review.lines[0].line_status, "Draft Created")
